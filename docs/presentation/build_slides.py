"""Build the 8-minute AMDockVS talk (English) as a .pptx.

Run:  PYTHONPATH=/tmp/pptxlib python3 docs/presentation/build_slides.py
Import the result into Google Slides (File > Import slides) or open in PowerPoint.
Dashed boxes are screenshot slots: replace them with real captures.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
LOGO = HERE.parents[1] / "src/amdockvs/ui/resources/icons/logo.png"
OUT = HERE / "AMDockVS_8min.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x41)
TEAL = RGBColor(0x0F, 0x8B, 0x8D)
ORANGE = RGBColor(0xE0, 0x7A, 0x1F)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xEE, 0xF3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP):
    """lines: str or list of str; a line starting with '• ' becomes a bullet."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.alignment = line, align
        p.space_after = Pt(8)
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size), color, bold
    return tb


def box(slide, x, y, w, h, label, fill=TEAL, fg=WHITE, size=14, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.alignment = line, PP_ALIGN.CENTER
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size if i == 0 else size - 3), fg, i == 0
    return s


def shot(slide, x, y, w, h, what):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = LIGHT
    s.line.color.rgb, s.line.dash_style, s.line.width = GREY, MSO_LINE_DASH_STYLE.DASH, Pt(1.5)
    tf = s.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text, p.alignment = f"[ CAPTURA ]\n{what}", PP_ALIGN.CENTER
    p.font.size, p.font.color.rgb = Pt(16), GREY


def slide(title, notes, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if kicker:
        text(s, 0.6, 0.35, 12, 0.4, kicker, size=13, color=TEAL, bold=True)
    text(s, 0.6, 0.65, 12, 0.9, title, size=32, bold=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def funnel(s, y, steps, fill):
    # ponytail: chevrons clip their own text, so the label is a textbox laid on top
    w = 12.1 / len(steps)
    for i, (head, sub) in enumerate(steps):
        x = 0.6 + i * w
        box(s, x, y, w, 1.25, "", fill=fill, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        tb = text(s, x + 0.3, y, w - 0.4, 1.25, [head, sub], size=11, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tf = tb.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0
        tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(14), True
        tf.paragraphs[0].space_after = Pt(2)


STEPS = ["Library", "Filtering", "Diversity", "Preparation", "Site", "Docking", "Analysis"]

# 1 — Title --------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
if LOGO.exists():
    s.shapes.add_picture(str(LOGO), Inches(0.9), Inches(2.3), height=Inches(2.6))
text(s, 4.0, 2.2, 8.8, 1.2, "AMDockVS", size=60, color=WHITE, bold=True)
text(s, 4.0, 3.4, 8.8, 1.2, "Integrated, reproducible, open-source virtual screening",
     size=26, color=RGBColor(0x9F, 0xE0, 0xE0))
text(s, 4.0, 5.2, 8.8, 0.9, ["Mario S. Valdés · [institution] · [event, date]"], size=16,
     color=WHITE)
s.notes_slide.notes_text_frame.text = (
    "(10 s) Introduce yourself. Today: the problem AMDockVS solves and what it can do.")

# 2 — Typical VS workflow -------------------------------------------------------
s = slide("A Typical Virtual Screening Workflow", kicker="THE PROBLEM", notes=(
    "(1:00) A virtual-screening campaign starts with libraries of millions, sometimes billions, "
    "of compounds (ZINC, Enamine REAL). We filter by properties and alerts, reduce the set by "
    "diversity, prepare the ligands and receptor, define the binding site, dock, and analyze. "
    "Each step often uses a different program and file format. The challenges are connecting "
    "tools by hand, converting formats, scaling to large libraries, and reproducing the exact "
    "workflow."))
funnel(s, 1.8, list(zip(STEPS, ["10⁶-10⁹ compounds", "Lipinski, PAINS", "clustering", "H⁺, 3D, PDBQT",
                                "pockets", "scoring", "poses, ranking"])), NAVY)
tools = ["ZINC\nEnamine", "RDKit\nKNIME", "RDKit\nscripts", "OpenBabel\nreduce · Meeko",
         "fpocket\nP2Rank", "Vina\nAutoDock", "PyMOL\nPLIP"]
for i, t in enumerate(tools):
    text(s, 0.6 + i * 12.1 / 7, 3.15, 12.1 / 7, 0.9, t.split("\n"), size=13, color=GREY,
         align=PP_ALIGN.CENTER)
for i, (h, d) in enumerate([("Fragmented", "one program and format per step"),
                            ("Manual", "ad hoc scripts and conversions"),
                            ("Scale", "from dozens to millions of molecules"),
                            ("Reproducibility", "which parameters were used?")]):
    box(s, 0.6 + i * 3.1, 4.8, 2.85, 1.6, f"{h}\n{d}", fill=ORANGE, size=18)

# 3 — Current solutions --------------------------------------------------------
s = slide("Current Solutions", kicker="CONTEXT", notes=(
    "(0:45) We are building on existing work. AMDock (2020) already offered a Vina and "
    "AutoDock4 docking GUI, but it handled one complex at a time. Other options include PyRx, "
    "web servers, VirtualFlow for HPC, and commercial suites. Each covers part of the need: "
    "ease of use, scale, or openness. A desktop tool combining a GUI, scale, openness, and "
    "reproducibility was missing."))
rows = [("", "GUI", "Scale", "Open source", "Reproducible"),
        ("AMDock 1.6.5", "✔", "✘  1 complex", "✔", "~"),
        ("PyRx", "✔", "~", "~", "~"),
        ("Web servers (SwissDock, DockThor)", "✔", "✘  quotas", "~", "~"),
        ("VirtualFlow", "✘  CLI / HPC", "✔", "✔", "✔"),
        ("Commercial suites (Schrödinger)", "✔", "✔", "✘  license", "✔"),
        ("AMDockVS", "✔", "✔", "✔  MIT", "✔")]
tbl = s.shapes.add_table(len(rows), 5, Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.6)).table
tbl.columns[0].width = Inches(4.5)
for c in range(1, 5):
    tbl.columns[c].width = Inches(1.9)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size, p.font.bold = Pt(17), r in (0, len(rows) - 1)
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else TEAL if r == len(rows) - 1 else \
            (LIGHT if r % 2 else WHITE)
        p.font.color.rgb = WHITE if r in (0, len(rows) - 1) else NAVY
text(s, 0.6, 6.6, 12, 0.5, "Valdés-Tresanco M.S. et al. AMDock. Biology Direct 15, 12 (2020).",
     size=11, color=GREY)

# 4 — AMDockVS -----------------------------------------------------------------
s = slide("AMDockVS: One Workflow, One Application", kicker="THE PROPOSAL", notes=(
    "(0:50) AMDockVS brings the entire workflow into one application. It offers two ways to "
    "use the same scientific logic: a desktop interface and a Python API for reproducible "
    "notebooks and scripts. It is MIT-licensed and installable with pip."))
funnel(s, 1.8, list(zip(STEPS, ["SDF · SMILES · PDB", "descriptors", "BitBIRCH", "RDKit · Meeko",
                                "P2Rank · manual", "Vina · Vinardo", "tables · PyMOL"])), TEAL)
box(s, 0.6, 3.5, 5.9, 2.9, "Desktop interface\nPySide6 + PyMOL · projects · tables · "
    "background jobs", fill=NAVY, size=22)
code = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(3.5), Inches(5.9), Inches(2.9))
code.fill.solid()
code.fill.fore_color.rgb = RGBColor(0x10, 0x18, 0x24)
code.line.fill.background()
tf = code.text_frame
tf.word_wrap = True
for i, line in enumerate(["# Python API — notebooks and scripts",
                          "from amdockvs import AMDockVSRuntime",
                          "rt = AMDockVSRuntime(...)",
                          "rt.loader ...      # import",
                          "rt.pockets.predict(...)",
                          "rt.docking ...     # prepare and dock"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.font.name, p.font.size = "Consolas", Pt(16)
    p.font.color.rgb = GREY if line.startswith("#") else RGBColor(0x9F, 0xE0, 0xE0)
text(s, 0.6, 6.55, 12.1, 0.5, "pip install AMDockVS   ·   Python 3.12   ·   MIT license",
     size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

# 5-8 — capability gallery -----------------------------------------------------
GALLERY = [
    ("Prepare and Select", "CAPABILITIES · 1/4",
     ["• Ligand and receptor tables with descriptors",
      "• Standardization, protonation, 3D generation, and minimization",
      "• Diversity clustering (BitBIRCH) → representative centroids"],
     "Ligand table + Prepare tab / Diversity results",
     "(0:45) Everything is available in filterable tables. Preparation handles protonation, "
     "3D generation, and minimization in one step. To reduce a large library, we use BitBIRCH "
     "clustering and keep the centroids: fewer compounds with the same diversity."),
    ("Binding Site and Docking", "CAPABILITIES · 2/4",
     ["• Pocket prediction with P2Rank or a manual box",
      "• Automatic preparation with Meeko",
      "• Parallel Vina / Vinardo runs; results sorted by affinity"],
     "P2Rank pockets on the receptor + affinity-ranked results table",
     "(0:45) If the binding site is unknown, P2Rank predicts pockets and we generate a box "
     "from them. Meeko prepares the PDBQT files, and docking runs as a parallel job. Results "
     "appear in a table that can be sorted by affinity."),
    ("Validate: Redocking", "CAPABILITIES · 3/4",
     ["• Redock the crystallographic ligand",
      "• RMSD between the predicted and experimental poses",
      "• Validate the protocol before screening"],
     "Redocking RMSD plot",
     "(0:45) Before screening, validate the protocol: redock the crystal ligand and measure "
     "the RMSD. [If data are available: 'the RMSD was below 2 Å for X of Y complexes.'] This "
     "shows whether the protocol reproduces the binding mode."),
    ("Explore: Off-Target and Analysis", "CAPABILITIES · 4/4",
     ["• One ligand against multiple receptors (off-target)",
      "• Aligned poses in PyMOL",
      "• 2D interaction diagrams · basic QSAR"],
     "Off-target view + PyMOL pose + interaction diagram",
     "(0:45) The off-target view compares one ligand against several receptors with aligned "
     "poses, which helps assess selectivity. Inspect each pose in PyMOL and with its 2D "
     "interaction diagram. The application also includes descriptors and basic QSAR models."),
]
for title, kicker, bullets, what, notes in GALLERY:
    s = slide(title, notes, kicker=kicker)
    shot(s, 0.6, 1.7, 7.9, 5.3, what)
    text(s, 8.8, 1.9, 4.0, 5.0, bullets, size=20)

# 9 — Architecture / ms_flow ---------------------------------------------------
s = slide("Architecture", kicker="UNDER THE HOOD", notes=(
    "(0:40) The architecture has three layers. At the top, the GUI and API use the same "
    "logic. AMDockVS provides the chemistry layer. At the bottom, ms_flow is our orchestrator: "
    "it manages projects, asynchronous jobs, executors, and CPU/GPU parallelism without a "
    "server or broker. ms_flow is general-purpose and supports large data workflows beyond chemistry."))
box(s, 0.6, 1.8, 5.9, 1.1, "Desktop interface\nPySide6 · PyMOL", fill=NAVY, size=20)
box(s, 6.8, 1.8, 5.9, 1.1, "Python API\nnotebooks · scripts", fill=NAVY, size=20)
box(s, 0.6, 3.1, 12.1, 1.2, "AMDockVS — scientific workflow\nRDKit · Meeko · Vina · P2Rank",
    fill=TEAL, size=22)
box(s, 0.6, 4.5, 12.1, 1.6, "ms_flow — orchestrator\nprojects · asynchronous jobs · "
    "executors · CPU/GPU parallelism · persistence", fill=ORANGE, size=24)
text(s, 0.6, 6.3, 12.1, 0.8, "ms_flow: no server or broker; reusable beyond chemistry "
     "—  pip install ms_flow", size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

# 10 — Conclusion --------------------------------------------------------------
s = slide("Conclusion", kicker="WRAP-UP", notes=(
    "(0:40) In summary: AMDockVS is integrated, reproducible, and works through either the "
    "GUI or code. Next steps include GNINA and GPU support, remote/HPC execution, and more "
    "docking engines. Close with the final statement."))
for i, (h, d) in enumerate([("Integrated", "the full workflow in one application"),
                            ("Reproducible", "projects, parameters, and API"),
                            ("Accessible", "GUI or Python · MIT · pip")]):
    box(s, 0.6 + i * 4.1, 1.8, 3.8, 1.6, f"{h}\n{d}", fill=TEAL, size=24)
text(s, 0.6, 3.7, 12.1, 0.5, "Next steps:  GNINA / GPU  ·  remote execution (HPC)  ·  "
     "more engines", size=18, color=GREY, align=PP_ALIGN.CENTER)
text(s, 1.2, 4.5, 10.9, 1.6, "“AMDockVS turns a fragmented set of docking tools into an "
     "integrated, reproducible, and accessible workflow for drug discovery.”",
     size=24, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
shot(s, 5.9, 6.1, 1.4, 1.2, "QR")
text(s, 7.5, 6.35, 5.2, 0.6, "GitHub · PyPI · docs  —  Thank you!", size=16, color=GREY)

prs.save(OUT)
print(OUT)
