"""Build the 8-minute AMDockVS talk (Spanish) as a .pptx.

Run:  PYTHONPATH=/tmp/pptxlib python3 docs/presentation/build_slides.py
Import the result into Google Slides (File > Import slides) or open in PowerPoint.
Dashed boxes are screenshot slots: replace them with real captures.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
LOGO = HERE.parents[1] / "src/amdockvs/ui/resources/icons/logo.png"
OUT = HERE / "AMDockVS_8min.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x41)
TEAL = RGBColor(0x0F, 0x8B, 0x8D)
ORANGE = RGBColor(0xE0, 0x7A, 0x1F)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xEE, 0xF3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP):
    """lines: str or list of str; a line starting with '• ' becomes a bullet."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.alignment = line, align
        p.space_after = Pt(8)
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size), color, bold
    return tb


def box(slide, x, y, w, h, label, fill=TEAL, fg=WHITE, size=14, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.alignment = line, PP_ALIGN.CENTER
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size if i == 0 else size - 3), fg, i == 0
    return s


def shot(slide, x, y, w, h, what):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = LIGHT
    s.line.color.rgb, s.line.dash_style, s.line.width = GREY, MSO_LINE_DASH_STYLE.DASH, Pt(1.5)
    tf = s.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text, p.alignment = f"[ CAPTURA ]\n{what}", PP_ALIGN.CENTER
    p.font.size, p.font.color.rgb = Pt(16), GREY


def slide(title, notes, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if kicker:
        text(s, 0.6, 0.35, 12, 0.4, kicker, size=13, color=TEAL, bold=True)
    text(s, 0.6, 0.65, 12, 0.9, title, size=32, bold=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def funnel(s, y, steps, fill):
    # ponytail: chevrons clip their own text, so the label is a textbox laid on top
    w = 12.1 / len(steps)
    for i, (head, sub) in enumerate(steps):
        x = 0.6 + i * w
        box(s, x, y, w, 1.25, "", fill=fill, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        tb = text(s, x + 0.3, y, w - 0.4, 1.25, [head, sub], size=11, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tf = tb.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0
        tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(14), True
        tf.paragraphs[0].space_after = Pt(2)


STEPS = ["Librería", "Filtrado", "Diversidad", "Preparación", "Sitio", "Docking", "Análisis"]

# 1 — Title --------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
if LOGO.exists():
    s.shapes.add_picture(str(LOGO), Inches(0.9), Inches(2.3), height=Inches(2.6))
text(s, 4.0, 2.2, 8.8, 1.2, "AMDockVS", size=60, color=WHITE, bold=True)
text(s, 4.0, 3.4, 8.8, 1.2, "Cribado virtual integrado, reproducible y de código abierto",
     size=26, color=RGBColor(0x9F, 0xE0, 0xE0))
text(s, 4.0, 5.2, 8.8, 0.9, ["Mario S. Valdés · [institución] · [evento, fecha]"], size=16,
     color=WHITE)
s.notes_slide.notes_text_frame.text = (
    "(10 s) Presentarse. Hoy: qué problema resuelve AMDockVS y qué puede hacer.")

# 2 — Typical VS workflow -------------------------------------------------------
s = slide("Un cribado virtual típico", kicker="EL PROBLEMA", notes=(
    "(1:00) Así es un cribado virtual: partimos de librerías de millones —incluso miles de "
    "millones— de compuestos (ZINC, Enamine REAL). Filtramos por propiedades y alertas, "
    "reducimos por diversidad, preparamos ligandos y receptor, definimos el sitio, hacemos "
    "docking y analizamos. Cada paso suele ser un programa distinto, con su formato. "
    "Los retos: pegar herramientas a mano, convertir formatos, escalar a muchos compuestos "
    "y poder repetir exactamente lo que se hizo."))
funnel(s, 1.8, list(zip(STEPS, ["10⁶–10⁹ cpds", "Lipinski, PAINS", "clustering", "H⁺, 3D, PDBQT",
                                "cavidades", "scoring", "poses, ranking"])), NAVY)
tools = ["ZINC\nEnamine", "RDKit\nKNIME", "RDKit\nscripts", "OpenBabel\nreduce · Meeko",
         "fpocket\nP2Rank", "Vina\nAutoDock", "PyMOL\nPLIP"]
for i, t in enumerate(tools):
    text(s, 0.6 + i * 12.1 / 7, 3.15, 12.1 / 7, 0.9, t.split("\n"), size=13, color=GREY,
         align=PP_ALIGN.CENTER)
for i, (h, d) in enumerate([("Fragmentado", "un programa y un formato por paso"),
                            ("Manual", "scripts y conversiones ad hoc"),
                            ("Escala", "de decenas a millones de moléculas"),
                            ("Reproducibilidad", "¿qué parámetros se usaron?")]):
    box(s, 0.6 + i * 3.1, 4.8, 2.85, 1.6, f"{h}\n{d}", fill=ORANGE, size=18)

# 3 — Current solutions --------------------------------------------------------
s = slide("Soluciones actuales", kicker="CONTEXTO", notes=(
    "(0:45) No partimos de cero. AMDock (2020) ya ofrecía una GUI para docking con Vina y "
    "AutoDock4, pero orientada a un complejo a la vez. Hay otras opciones: PyRx, servidores "
    "web, VirtualFlow para HPC, suites comerciales. Cada una cubre una parte: o es fácil de "
    "usar, o escala, o es abierta. Faltaba algo intermedio: GUI + escala + abierto + "
    "reproducible, en el escritorio."))
rows = [("", "GUI", "Escala", "Abierto", "Reproducible"),
        ("AMDock 1.6.5", "✔", "✘  1 complejo", "✔", "~"),
        ("PyRx", "✔", "~", "~", "~"),
        ("Servidores web (SwissDock, DockThor)", "✔", "✘  cuotas", "~", "~"),
        ("VirtualFlow", "✘  CLI / HPC", "✔", "✔", "✔"),
        ("Suites comerciales (Schrödinger)", "✔", "✔", "✘  licencia", "✔"),
        ("AMDockVS", "✔", "✔", "✔  MIT", "✔")]
tbl = s.shapes.add_table(len(rows), 5, Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.6)).table
tbl.columns[0].width = Inches(4.5)
for c in range(1, 5):
    tbl.columns[c].width = Inches(1.9)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size, p.font.bold = Pt(17), r in (0, len(rows) - 1)
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else TEAL if r == len(rows) - 1 else \
            (LIGHT if r % 2 else WHITE)
        p.font.color.rgb = WHITE if r in (0, len(rows) - 1) else NAVY
text(s, 0.6, 6.6, 12, 0.5, "Valdés-Tresanco M.S. et al. AMDock. Biology Direct 15, 12 (2020).",
     size=11, color=GREY)

# 4 — AMDockVS -----------------------------------------------------------------
s = slide("AMDockVS: un flujo, una aplicación", kicker="LA PROPUESTA", notes=(
    "(0:50) AMDockVS integra todo el embudo en una sola aplicación. Se usa de dos formas "
    "sobre la misma lógica científica: una interfaz de escritorio, o una API de Python para "
    "notebooks y scripts reproducibles. Es MIT y se instala con pip."))
funnel(s, 1.8, list(zip(STEPS, ["SDF · SMILES · PDB", "descriptores", "BitBIRCH", "RDKit · Meeko",
                                "P2Rank · manual", "Vina · Vinardo", "tablas · PyMOL"])), TEAL)
box(s, 0.6, 3.5, 5.9, 2.9, "Interfaz de escritorio\nPySide6 + PyMOL · proyectos · tablas · "
    "trabajos en segundo plano", fill=NAVY, size=22)
code = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(3.5), Inches(5.9), Inches(2.9))
code.fill.solid()
code.fill.fore_color.rgb = RGBColor(0x10, 0x18, 0x24)
code.line.fill.background()
tf = code.text_frame
tf.word_wrap = True
for i, line in enumerate(["# API de Python — notebooks y scripts",
                          "from amdockvs import AMDockVSRuntime",
                          "rt = AMDockVSRuntime(...)",
                          "rt.loader ...      # importar",
                          "rt.pockets.predict(...)",
                          "rt.docking ...     # preparar y acoplar"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.font.name, p.font.size = "Consolas", Pt(16)
    p.font.color.rgb = GREY if line.startswith("#") else RGBColor(0x9F, 0xE0, 0xE0)
text(s, 0.6, 6.55, 12.1, 0.5, "pip install AMDockVS   ·   Python 3.12   ·   licencia MIT",
     size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

# 5-8 — capability gallery -----------------------------------------------------
GALLERY = [
    ("Preparar y seleccionar", "CAPACIDADES · 1/4",
     ["• Tablas de ligandos y receptores con descriptores",
      "• Estandarización, protonación, 3D y minimización",
      "• Clustering por diversidad (BitBIRCH) → centroides representativos"],
     "Tabla de ligandos + pestaña Prepare / resultado de Diversity",
     "(0:45) Todo se ve en tablas filtrables. La preparación hace protonación, 3D y "
     "minimización en un solo paso. Para reducir una librería grande usamos clustering "
     "BitBIRCH y nos quedamos con los centroides: menos compuestos, misma diversidad."),
    ("Sitio de unión y docking", "CAPACIDADES · 2/4",
     ["• Predicción de cavidades con P2Rank o caja manual",
      "• Preparación automática con Meeko",
      "• Vina / Vinardo en paralelo; resultados ordenados por afinidad"],
     "Cavidades P2Rank sobre el receptor + tabla de resultados por afinidad",
     "(0:45) Si no conocemos el sitio, P2Rank predice las cavidades y generamos la caja a "
     "partir de ellas. Meeko prepara los PDBQT y el docking se lanza como trabajo en paralelo. "
     "Los resultados salen en una tabla ordenable por afinidad."),
    ("Validar: redocking", "CAPACIDADES · 3/4",
     ["• Re-acoplar el ligando cristalográfico",
      "• RMSD entre pose predicha y experimental",
      "• Confirma el protocolo antes del cribado"],
     "Gráfica de RMSD del redocking",
     "(0:45) Antes de cribar, validamos: re-acoplamos el ligando del cristal y medimos el RMSD. "
     "[Si hay datos: 'en X de Y complejos el RMSD fue < 2 Å'.] Así sabemos que el protocolo "
     "reproduce el modo de unión."),
    ("Explorar: off-target y análisis", "CAPACIDADES · 4/4",
     ["• Un ligando frente a varios receptores (off-target)",
      "• Poses alineadas en PyMOL",
      "• Diagramas 2D de interacciones · QSAR básico"],
     "Vista off-target + pose en PyMOL + diagrama de interacciones",
     "(0:45) La vista off-target pone un ligando frente a varios receptores con las poses "
     "alineadas: útil para selectividad. Cada pose se inspecciona en PyMOL y con su diagrama "
     "2D de interacciones. También hay descriptores y modelos QSAR básicos."),
]
for title, kicker, bullets, what, notes in GALLERY:
    s = slide(title, notes, kicker=kicker)
    shot(s, 0.6, 1.7, 7.9, 5.3, what)
    text(s, 8.8, 1.9, 4.0, 5.0, bullets, size=20)

# 9 — Architecture / ms_flow ---------------------------------------------------
s = slide("Arquitectura", kicker="POR DENTRO", notes=(
    "(0:40) Por dentro hay tres capas. Arriba, la GUI y la API usan la misma lógica. "
    "En medio, AMDockVS: la química. Abajo, ms_flow: nuestro orquestador, que gestiona "
    "proyectos, trabajos asíncronos, ejecutores y paralelismo CPU/GPU, sin servidor ni broker. "
    "ms_flow es genérico: sirve para cualquier flujo masivo de datos, no solo química."))
box(s, 0.6, 1.8, 5.9, 1.1, "Interfaz de escritorio\nPySide6 · PyMOL", fill=NAVY, size=20)
box(s, 6.8, 1.8, 5.9, 1.1, "API de Python\nnotebooks · scripts", fill=NAVY, size=20)
box(s, 0.6, 3.1, 12.1, 1.2, "AMDockVS — lógica científica\nRDKit · Meeko · Vina · P2Rank",
    fill=TEAL, size=22)
box(s, 0.6, 4.5, 12.1, 1.6, "ms_flow — orquestador\nproyectos · trabajos asíncronos · "
    "ejecutores · paralelismo CPU/GPU · persistencia", fill=ORANGE, size=24)
text(s, 0.6, 6.3, 12.1, 0.8, "ms_flow: sin servidor ni broker, reutilizable fuera de la química "
     "—  pip install ms_flow", size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

# 10 — Conclusion --------------------------------------------------------------
s = slide("Conclusión", kicker="CIERRE", notes=(
    "(0:40) Resumen: integra, es reproducible y sirve desde la GUI o desde código. "
    "Próximos pasos: GNINA y GPU, ejecución remota/HPC y más motores. Frase final."))
for i, (h, d) in enumerate([("Integra", "todo el embudo en una aplicación"),
                            ("Reproducible", "proyectos, parámetros y API"),
                            ("Accesible", "GUI o Python · MIT · pip")]):
    box(s, 0.6 + i * 4.1, 1.8, 3.8, 1.6, f"{h}\n{d}", fill=TEAL, size=24)
text(s, 0.6, 3.7, 12.1, 0.5, "Próximos pasos:  GNINA / GPU  ·  ejecución remota (HPC)  ·  "
     "nuevos motores", size=18, color=GREY, align=PP_ALIGN.CENTER)
text(s, 1.2, 4.5, 10.9, 1.6, "“AMDockVS convierte un conjunto fragmentado de herramientas de "
     "docking en un flujo integrado, reproducible y accesible para el diseño de fármacos.”",
     size=24, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
shot(s, 5.9, 6.1, 1.4, 1.2, "QR")
text(s, 7.5, 6.35, 5.2, 0.6, "github · pypi · docs  —  ¡Gracias!", size=16, color=GREY)

prs.save(OUT)
print(OUT)
